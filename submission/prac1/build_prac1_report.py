from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "PRAC1_Report_Source.md"
OUTPUT = ROOT / "IWFC_PRAC1_Report.docx"
ASSETS = ROOT / "assets"
ASSETS.mkdir(parents=True, exist_ok=True)

INK = "17202A"
NAVY = "193A5A"
BLUE = "2B6F9E"
PALE = "EAF2F8"
MID = "D9E2EA"
WHITE = "FFFFFF"
GREY = "5F6B73"


def font(size: int, bold: bool = False):
    choices = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for choice in choices:
        if choice.exists():
            return ImageFont.truetype(str(choice), size)
    return ImageFont.load_default()


def rounded_box(draw, xy, title, lines, fill, outline=NAVY, title_size=30, body_size=22):
    x1, y1, x2, y2 = xy
    fill_color = fill if fill.startswith("#") else f"#{fill}"
    draw.rounded_rectangle(xy, radius=18, fill=fill_color, outline=f"#{outline}", width=4)
    draw.text((x1 + 18, y1 + 14), title, font=font(title_size, True), fill=f"#{INK}")
    y = y1 + 58
    for line in lines:
        draw.text((x1 + 20, y), line, font=font(body_size), fill=f"#{INK}")
        y += body_size + 10


def arrow(draw, start, end, color=NAVY, width=5):
    draw.line([start, end], fill=f"#{color}", width=width)
    x2, y2 = end
    x1, y1 = start
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    size = 16
    points = [
        (x2, y2),
        (x2 - ux * size + px * size * 0.55, y2 - uy * size + py * size * 0.55),
        (x2 - ux * size - px * size * 0.55, y2 - uy * size - py * size * 0.55),
    ]
    draw.polygon(points, fill=f"#{color}")


def create_uml():
    image = Image.new("RGB", (2200, 1580), "white")
    draw = ImageDraw.Draw(image)
    draw.text((60, 35), "IWFC implemented class architecture", font=font(48, True), fill=f"#{INK}")

    rounded_box(draw, (760, 120, 1440, 300), "IWFCFacade  <<Facade>>",
                ["public use-case API", "role and active-account guards", "console entry point"], PALE)
    rounded_box(draw, (170, 390, 760, 710), "Domain model", [
        "User  <<abstract>>", "Administrator extends User", "Instructor extends User",
        "Member extends User", "Equipment", "FitnessSession", "MaintenanceRequest"
    ], "F7F9FB")
    rounded_box(draw, (820, 390, 1380, 610), "Application services", [
        "BookingService", "MaintenanceService  <<Subject>>", "User observers receive events"
    ], "F7F9FB")
    rounded_box(draw, (1440, 390, 2030, 610), "Creation and storage", [
        "EntityFactory  <<Factory>>", "GenericRepository<T>", "HashMap<String,T> storage"
    ], "F7F9FB")
    rounded_box(draw, (470, 820, 1110, 1040), "Checked custom exceptions", [
        "InvalidBookingException", "UnauthorizedAccessException", "DuplicateDataException"
    ], "FFF7E6")
    rounded_box(draw, (1160, 820, 1900, 1040), "Key relationships", [
        "Session -> Instructor, Equipment, Members", "Request -> Equipment and reporting Instructor",
        "Facade -> Factory, services, repositories"
    ], "FFF7E6")
    rounded_box(draw, (410, 1190, 1790, 1450), "Architectural rule", [
        "Console -> public Facade -> services -> domain and typed repositories",
        "15 top-level production files; nested enums remain within their owners",
        "No presentation code mutates a repository or service directly"
    ], PALE, title_size=32, body_size=25)

    arrow(draw, (1100, 300), (1100, 380))
    arrow(draw, (900, 300), (600, 380))
    arrow(draw, (1300, 300), (1700, 380))
    arrow(draw, (600, 710), (720, 810))
    arrow(draw, (1100, 610), (920, 810))
    arrow(draw, (1700, 610), (1550, 810))
    arrow(draw, (790, 1040), (900, 1180))
    arrow(draw, (1530, 1040), (1370, 1180))

    image.save(ASSETS / "figure-1-uml-architecture.png", quality=95)


def create_maintenance_flow():
    image = Image.new("RGB", (2100, 980), "white")
    draw = ImageDraw.Draw(image)
    draw.text((55, 35), "Maintenance workflow and Observer delivery", font=font(46, True), fill=f"#{INK}")

    boxes = [
        ((80, 220, 600, 480), "PENDING", ["Instructor reports fault", "Equipment -> FAULTY", "Notify active Administrators"]),
        ((790, 220, 1310, 480), "ASSIGNED", ["Administrator assigns", "Equipment -> UNDER MAINTENANCE", "Notify reporting Instructor"]),
        ((1500, 220, 2020, 480), "COMPLETED", ["Administrator completes", "Active equipment -> OPERATIONAL", "Notify reporting Instructor"]),
    ]
    for xy, title, lines in boxes:
        rounded_box(draw, xy, title, lines, PALE, title_size=34, body_size=25)
    arrow(draw, (610, 350), (780, 350), BLUE, 7)
    arrow(draw, (1320, 350), (1490, 350), BLUE, 7)
    draw.text((635, 300), "assign", font=font(24, True), fill=f"#{BLUE}")
    draw.text((1345, 300), "complete", font=font(24, True), fill=f"#{BLUE}")

    rounded_box(draw, (330, 650, 1770, 880), "Preventative cycle", [
        "Alert at >= 100 cumulative hours -> one message per active Administrator",
        "Repeated checks in the same cycle do not duplicate the alert",
        "Completion clears the latch and sets next due = current cumulative usage + 100 hours"
    ], "FFF7E6", title_size=32, body_size=26)
    image.save(ASSETS / "figure-2-maintenance-workflow.png", quality=95)


def create_test_chart():
    image = Image.new("RGB", (1900, 1050), "white")
    draw = ImageDraw.Draw(image)
    draw.text((55, 35), "Passing tests by architectural area", font=font(46, True), fill=f"#{INK}")
    data = [("Facade and console", 46), ("Domain entities", 34), ("Services", 33), ("Factory and repository", 13)]
    max_value = 50
    x0, y0, bar_max, bar_h, gap = 410, 190, 1250, 120, 78
    for index, (label, value) in enumerate(data):
        y = y0 + index * (bar_h + gap)
        draw.text((55, y + 34), label, font=font(27, True), fill=f"#{INK}")
        draw.rounded_rectangle((x0, y, x0 + bar_max, y + bar_h), radius=18, fill="#E8EDF1")
        width = int(bar_max * value / max_value)
        draw.rounded_rectangle((x0, y, x0 + width, y + bar_h), radius=18, fill=f"#{BLUE}")
        draw.text((x0 + width + 20, y + 32), str(value), font=font(30, True), fill=f"#{INK}")
    draw.text((55, 940), "Total: 126 tests | Failures: 0 | Errors: 0 | Skipped: 0",
              font=font(30, True), fill=f"#{NAVY}")
    image.save(ASSETS / "figure-3-test-distribution.png", quality=95)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Arial"
    run.font.size = Pt(9)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def style_run(run, name="Arial", size=12, bold=False, italic=False, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text):
    token_re = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")
    position = 0
    for match in token_re.finditer(text):
        if match.start() > position:
            style_run(paragraph.add_run(text[position:match.start()]))
        token = match.group(0)
        if token.startswith("`"):
            style_run(paragraph.add_run(token[1:-1]), "Consolas", 10.5, color=NAVY)
        elif token.startswith("**"):
            style_run(paragraph.add_run(token[2:-2]), bold=True)
        else:
            style_run(paragraph.add_run(token[1:-1]), italic=True)
        position = match.end()
    if position < len(text):
        style_run(paragraph.add_run(text[position:]))


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    title.paragraph_format.space_after = Pt(18)

    for style_name, size in (("Heading 1", 17), ("Heading 2", 14)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14 if style_name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(7)

    caption = doc.styles["Caption"]
    caption.font.name = "Arial"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    caption.font.size = Pt(10)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(GREY)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = False
    caption.paragraph_format.space_after = Pt(9)

    if "Code Block" not in doc.styles:
        code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        code.font.size = Pt(9)
        code.font.color.rgb = RGBColor.from_string(INK)
        code.paragraph_format.left_indent = Cm(0.7)
        code.paragraph_format.space_after = Pt(3)


def add_cover(doc):
    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(85)
    paragraph.add_run("Intelligent Wellness and Fitness Center Java Prototype")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(subtitle.add_run("CMP 7001 Advanced Programming\nPRAC1 Practical Project Report"),
              size=15, bold=True, color=NAVY)
    subtitle.paragraph_format.space_after = Pt(48)

    details = [
        ("Student", "Mohamed Rifad"),
        ("Student ID", "[STUDENT ID]"),
        ("Module tutor", "[MODULE TUTOR]"),
        ("Submission date", "[SUBMISSION DATE]"),
        ("Approximate report word count", "2,934 words before references and appendix"),
    ]
    table = doc.add_table(rows=len(details), cols=2)
    table.alignment = 1
    table.autofit = False
    table.columns[0].width = Cm(5.0)
    table.columns[1].width = Cm(9.0)
    for index, (label, value) in enumerate(details):
        left, right = table.rows[index].cells
        set_cell_margins(left)
        set_cell_margins(right)
        style_run(left.paragraphs[0].add_run(label), bold=True, color=NAVY)
        style_run(right.paragraphs[0].add_run(value))
        for cell in (left, right):
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:color"), MID)
            borders.append(bottom)
            tc_pr.append(borders)

    repository = doc.add_paragraph()
    repository.alignment = WD_ALIGN_PARAGRAPH.CENTER
    repository.paragraph_format.space_before = Pt(40)
    style_run(repository.add_run("Private project repository\n"), size=10, bold=True, color=NAVY)
    style_run(repository.add_run("github.com/MohamedRifad/intelligent-wellness-and-fitness-center"),
              size=10, color=GREY)
    doc.add_page_break()


def add_contents(doc):
    doc.add_heading("Contents", level=1)
    items = [
        "Executive Summary", "1 Introduction", "2 Requirements Analysis and Solution Scope",
        "3 Solution Design and 15 Class Architecture", "4 Design Pattern Application",
        "5 Object Oriented Principles and Advanced Constructs",
        "6 Workflow Implementation and Rule Interaction", "7 Exception Handling and Robustness",
        "8 Verification and Test Results", "9 Critical Evaluation", "10 Conclusion", "References",
        "Appendix A Verification Commands"
    ]
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        style_run(paragraph.add_run(item), size=11)
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    style_run(note.add_run("Submission note: attach the institutional coversheet and feedback sheet as required by the brief."),
              size=10, italic=True, color=GREY)
    doc.add_page_break()


def build_document():
    create_uml()
    create_maintenance_flow()
    create_test_chart()

    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.gutter = Inches(0.5)
    section.header_distance = Inches(1.0)
    section.footer_distance = Inches(1.0)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(header.add_run("CMP 7001 PRAC1  |  IWFC"), size=9, color=GREY)
    add_page_number(section.footer.paragraphs[0])

    core = doc.core_properties
    core.title = "Intelligent Wellness and Fitness Center Java Prototype"
    core.subject = "CMP 7001 PRAC1 Practical Project Report"
    core.author = "Mohamed Rifad"
    core.keywords = "Java, object-oriented programming, design patterns, JUnit, IWFC"

    add_cover(doc)
    add_contents(doc)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = next(index for index, line in enumerate(lines) if line == "## Executive Summary")
    lines = lines[start:]
    paragraph_buffer = []
    in_code = False
    code_lines = []

    def flush_paragraph():
        if not paragraph_buffer:
            return
        text = " ".join(part.strip() for part in paragraph_buffer).strip()
        paragraph_buffer.clear()
        if not text:
            return
        if text.startswith("**Figure") and text.endswith("**"):
            paragraph = doc.add_paragraph(style="Caption")
            add_inline(paragraph, text[2:-2])
        else:
            paragraph = doc.add_paragraph()
            add_inline(paragraph, text)

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                for code_line in code_lines:
                    paragraph = doc.add_paragraph(style="Code Block")
                    paragraph.add_run(code_line)
                code_lines.clear()
                in_code = False
            else:
                flush_paragraph()
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not stripped:
            flush_paragraph()
            continue
        if stripped.startswith("## "):
            flush_paragraph()
            heading = stripped[3:]
            if heading in {"References", "Appendix A Verification Commands"}:
                doc.add_page_break()
            doc.add_heading(heading, level=1)
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            doc.add_heading(stripped[4:], level=2)
            continue
        image_match = re.fullmatch(r"!\[[^]]*]\(([^)]+)\)", stripped)
        if image_match:
            flush_paragraph()
            path = ROOT / image_match.group(1)
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(path), width=Inches(5.55))
            continue
        paragraph_buffer.append(stripped)
    flush_paragraph()

    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Heading 1":
            paragraph.paragraph_format.page_break_before = False
        if paragraph.text.startswith(("CMP 7001 (", "Fowler,", "Gamma,", "GitHub (", "JUnit Team",
                                      "Oracle (", "Rifad,")):
            paragraph.paragraph_format.left_indent = Cm(0.75)
            paragraph.paragraph_format.first_line_indent = Cm(-0.75)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_document()
